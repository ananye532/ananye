"""File bytes -> plain text. Supports PDF, DOCX and plain text."""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"
MIN_TEXT_CHARS = 80


class ExtractionError(ValueError):
    """Any file the extractor will not or cannot process. `code` is stable for API clients."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class Extracted:
    text: str
    source: str  # pdf | docx | txt | paste
    page_count: int | None


_BULLETS = "•●▪■‣⁃◦∙➢✓✔"


def clean_text(raw: str) -> str:
    text = unicodedata.normalize("NFKC", raw)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(f"[{_BULLETS}]", "•", text)
    text = re.sub(r"[​‌‍﻿]", "", text)
    text = re.sub(r"[  ]{2,}", " ", text)
    lines = [ln.strip() for ln in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes, max_pages: int) -> Extracted:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfplumber.utils.exceptions import PdfminerException

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            count = len(pdf.pages)
            if count == 0:
                raise ExtractionError("empty_file", "The PDF has no pages.")
            if count > max_pages:
                raise ExtractionError("too_many_pages", f"The PDF has {count} pages; the limit is {max_pages}.")
            pages = [p.extract_text(x_tolerance=1.5, y_tolerance=3) or "" for p in pdf.pages]
    except ExtractionError:
        raise
    except PDFPasswordIncorrect as exc:
        raise ExtractionError("encrypted", "Password-protected PDFs are not supported.") from exc
    except PdfminerException as exc:
        if exc.args and isinstance(exc.args[0], PDFPasswordIncorrect):
            raise ExtractionError("encrypted", "Password-protected PDFs are not supported.") from exc
        raise ExtractionError("corrupt", "The PDF could not be read.") from exc
    except Exception as exc:  # pdfminer raises many unrelated types on malformed input
        raise ExtractionError("corrupt", "The PDF could not be read.") from exc
    return Extracted(clean_text("\n".join(pages)), "pdf", count)


def _docx(data: bytes) -> Extracted:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("corrupt", "The Word document could not be read.") from exc
    parts: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        prefix = "• " if "list" in style and para.text.strip() else ""
        parts.append(prefix + para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return Extracted(clean_text("\n".join(parts)), "docx", None)


def _txt(data: bytes) -> Extracted:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return Extracted(clean_text(data.decode(enc)), "txt", None)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("unsupported", "The text file encoding is not supported.")


def extract(filename: str, data: bytes, *, max_pages: int = 10) -> Extracted:
    name = (filename or "").lower()
    if data.startswith(PDF_MAGIC):
        result = _pdf(data, max_pages)
    elif data.startswith(ZIP_MAGIC) and name.endswith(".docx"):
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ExtractionError("corrupt", "The Word document could not be read.")
        result = _docx(data)
    elif name.endswith((".txt", ".md")):
        result = _txt(data)
    else:
        raise ExtractionError("unsupported", "Upload a PDF, DOCX or TXT file.")
    if len(result.text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "no_text",
            "Almost no text could be extracted. If this is a scanned image, export a text-based PDF instead.",
        )
    return result


def from_paste(text: str) -> Extracted:
    cleaned = clean_text(text)
    if len(cleaned) < MIN_TEXT_CHARS:
        raise ExtractionError("no_text", "The pasted resume is too short to analyze.")
    return Extracted(cleaned, "paste", None)
